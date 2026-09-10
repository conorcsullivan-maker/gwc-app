"""Build the GWC front-page dashboard inside the official workbook.

Usage: python3 build_dashboard.py "path/to/workbook.xlsx"

Adds (or rebuilds) two sheets:
  - "Dashboard"  — the interactive front page (season selector drives it)
  - "GWC Calc"   — hidden formula engine the dashboard reads from

Everything is live Excel formulas over the raw data table on the 'Games'
sheet (rows 27-1000), so the dashboard updates itself as new picks are
synced or typed in. A timestamped backup is saved first.
"""
from __future__ import annotations
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ------------------------------------------------------------------- palette
INK = "0B0B0B"          # primary text
INK2 = "52514E"         # secondary text
MUTED = "898781"        # labels
GRID = "E1E0D9"         # hairlines
BASE = "C3C2B7"         # stronger rules
TILE = "F7F6F3"         # tile fill
BANNER = "1A1A19"       # banner fill
GOLD = "C9A227"         # brand accent (chrome only, never data)
BLUE = "2A78D6"         # sequential / bars
RED = "E34948"          # diverging warm pole
NEUTRAL = "F0EFEC"      # diverging midpoint
GOOD = "006300"         # success text (light surface)
CRIT = "D03B3B"         # critical
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]

MEMBERS = ["Conor", "Parker", "Grundy", "JR", "Jimmy", "Spencer"]
N = len(MEMBERS)
WEEKS = 18

# raw table on the Games sheet (header row 26, data 27+, room to grow)
R0, R1 = 27, 1000
G = {k: f"Games!${c}${R0}:${c}${R1}" for k, c in
     {"season": "A", "week": "B", "day": "C", "person": "D", "pick": "F",
      "result": "H", "ptype": "I", "favdog": "J", "homeaway": "K"}.items()}

SEASONS = ["All seasons", "2025 / 2026", "2026 / 2027"]
SEASON_CELL = "Dashboard!$R$3"
SCRIT = "'GWC Calc'!$B$2"


def wcol(w):
    return get_column_letter(2 + w)          # week 1 -> C ... week 18 -> T


# =========================================================== the calc engine
# layout (rows):  4 week numbers · 5-10 wins · 12-17 losses · 19-24 decided ·
# 26-31 pct · 33-38 cum wins · 40-45 cum losses · 47-52 cum pct (NA gaps) ·
# 54-57 coop weekly · 59-60 scalars · 61-67 per-player summary ·
# 70-75 sorted standings · 78-86 splits · 92-96 days · 99-102 awards
def build_calc(ws):
    ws["A1"] = "GWC dashboard engine — auto-generated, edit nothing here"
    ws["B1"] = f"={SEASON_CELL}"
    ws["B2"] = '=IF($B$1="All seasons","*",$B$1)'

    for w in range(1, WEEKS + 1):
        ws.cell(row=4, column=2 + w, value=w)

    for i, m in enumerate(MEMBERS):
        for base in (5, 12, 19, 26, 33, 40, 47):
            ws.cell(row=base + i, column=1, value=m)
        rw, rl, rd, rp = 5 + i, 12 + i, 19 + i, 26 + i
        rcw, rcl, rcp = 33 + i, 40 + i, 47 + i
        for w in range(1, WEEKS + 1):
            c = wcol(w)
            ws[f"{c}{rw}"] = (f'=COUNTIFS({G["person"]},$A{rw},{G["week"]},'
                              f'{c}$4,{G["result"]},"Win",{G["season"]},{SCRIT})')
            ws[f"{c}{rl}"] = (f'=COUNTIFS({G["person"]},$A{rl},{G["week"]},'
                              f'{c}$4,{G["result"]},"Loss",{G["season"]},{SCRIT})')
            ws[f"{c}{rd}"] = f"={c}{rw}+{c}{rl}"
            ws[f"{c}{rp}"] = f'=IF({c}{rd}=0,"",{c}{rw}/{c}{rd})'
            ws[f"{c}{rcw}"] = f"=SUM($C{rw}:{c}{rw})"
            ws[f"{c}{rcl}"] = f"=SUM($C{rl}:{c}{rl})"
            ws[f"{c}{rcp}"] = (f"=IF({c}{rcw}+{c}{rcl}=0,NA(),"
                               f"{c}{rcw}/({c}{rcw}+{c}{rcl}))")

    ws["A54"], ws["A55"] = "coop wins", "coop losses"
    ws["A56"], ws["A57"] = "coop pct", "week max pct"
    for w in range(1, WEEKS + 1):
        c = wcol(w)
        ws[f"{c}54"] = f"=SUM({c}5:{c}{4 + N})"
        ws[f"{c}55"] = f"=SUM({c}12:{c}{11 + N})"
        ws[f"{c}56"] = f'=IF({c}54+{c}55=0,"",{c}54/({c}54+{c}55))'
        ws[f"{c}57"] = f"=MAX({c}26:{c}{25 + N})"
    ws["A59"] = "last week with picks"
    ws["B59"] = "=SUMPRODUCT(MAX(((C54:T54+C55:T55)>0)*C4:T4))"
    ws["A60"] = "best coop week: pct, week#"
    ws["B60"] = "=IF(COUNT(C56:T56)=0,0,MAX(C56:T56))"
    ws["C60"] = "=IFERROR(MATCH(B60,C56:T56,0),0)"

    # per-player summary block, rows 62-67:
    # B name · C W · D L · E P · F picks · G pct · H sort score · I champs ·
    # J perfect · K goose · L l4 wins · M l4 losses · N l4 pct · O diff · P tag
    hdr = ["player", "W", "L", "P", "picks", "pct", "score", "champs",
           "perfect", "goose", "l4w", "l4l", "l4pct", "diff", "form"]
    for j, h in enumerate(hdr):
        ws.cell(row=61, column=2 + j, value=h)
    for i, m in enumerate(MEMBERS):
        r = 62 + i
        rw, rl, rd, rp = 5 + i, 12 + i, 19 + i, 26 + i
        ws[f"B{r}"] = m
        ws[f"C{r}"] = (f'=COUNTIFS({G["person"]},$B{r},{G["result"]},"Win",'
                       f'{G["season"]},{SCRIT})')
        ws[f"D{r}"] = (f'=COUNTIFS({G["person"]},$B{r},{G["result"]},"Loss",'
                       f'{G["season"]},{SCRIT})')
        ws[f"E{r}"] = (f'=COUNTIFS({G["person"]},$B{r},{G["result"]},"Push",'
                       f'{G["season"]},{SCRIT})')
        ws[f"F{r}"] = f"=C{r}+D{r}+E{r}"
        ws[f"G{r}"] = f"=IF(C{r}+D{r}=0,0,C{r}/(C{r}+D{r}))"
        ws[f"H{r}"] = f"=G{r}+C{r}*0.0001+{5 - i}*0.0000001"
        ws[f"I{r}"] = f"=SUMPRODUCT((C{rp}:T{rp}=C$57:T$57)*(C{rd}:T{rd}>0))"
        ws[f"J{r}"] = f"=SUMPRODUCT((C{rw}:T{rw}>=2)*(C{rl}:T{rl}=0))"
        ws[f"K{r}"] = f"=SUMPRODUCT((C{rw}:T{rw}=0)*(C{rl}:T{rl}>=2))"
        ws[f"L{r}"] = f"=SUMPRODUCT((C$4:T$4>$B$59-4)*C{rw}:T{rw})"
        ws[f"M{r}"] = f"=SUMPRODUCT((C$4:T$4>$B$59-4)*C{rl}:T{rl})"
        ws[f"N{r}"] = f'=IF(L{r}+M{r}=0,"",L{r}/(L{r}+M{r}))'
        ws[f"O{r}"] = f'=IF(N{r}="","",N{r}-G{r})'
        ws[f"P{r}"] = (f'=IF(O{r}="","—",IF(O{r}>=0.1,"▲ HOT",'
                       f'IF(O{r}<=-0.1,"▼ COLD","► STEADY")))')

    # sorted standings, rows 70-75:
    # B source-row index · C name · D W · E L · F P · G pct · H bar pct ·
    # I champs · J perfect · K goose · L l4 pct · M form tag
    src_cols = {"C": "B", "D": "C", "E": "D", "F": "E", "G": "G", "H": "G",
                "I": "I", "J": "J", "K": "K", "L": "N", "M": "P"}
    for k in range(N):
        r = 70 + k
        ws[f"B{r}"] = (f"=MATCH(LARGE($H$62:$H${61 + N},{k + 1}),"
                       f"$H$62:$H${61 + N},0)")
        for dst, src in src_cols.items():
            ws[f"{dst}{r}"] = f"=INDEX(${src}$62:${src}${61 + N},$B{r})"

    splits = [
        ("Spreads", f'{G["ptype"]},"Spread"'),
        ("Over/Unders", f'{G["ptype"]},"Over/Under"'),
        ("Moneylines", f'{G["ptype"]},"Moneyline"'),
        ("Favorites", f'{G["favdog"]},"Favorite"'),
        ("Underdogs", f'{G["favdog"]},"Underdog"'),
        ("Home teams", f'{G["homeaway"]},"Home"'),
        ("Away teams", f'{G["homeaway"]},"Away"'),
        ("Overs", f'{G["pick"]},"Over *"'),
        ("Unders", f'{G["pick"]},"Under *"'),
    ]
    for i, (label, cond) in enumerate(splits):
        r = 78 + i
        ws[f"B{r}"] = label
        ws[f"C{r}"] = f'=COUNTIFS({cond},{G["result"]},"Win",{G["season"]},{SCRIT})'
        ws[f"D{r}"] = f'=COUNTIFS({cond},{G["result"]},"Loss",{G["season"]},{SCRIT})'
        ws[f"E{r}"] = f"=IF(C{r}+D{r}=0,0,C{r}/(C{r}+D{r}))"

    for i, d in enumerate(["Thu", "Fri", "Sat", "Sun", "Mon"]):
        r = 92 + i
        ws[f"B{r}"] = d
        ws[f"C{r}"] = (f'=COUNTIFS({G["day"]},"{d}",{G["result"]},"Win",'
                       f'{G["season"]},{SCRIT})')
        ws[f"D{r}"] = (f'=COUNTIFS({G["day"]},"{d}",{G["result"]},"Loss",'
                       f'{G["season"]},{SCRIT})')
        ws[f"E{r}"] = f"=IF(C{r}+D{r}=0,0,C{r}/(C{r}+D{r}))"

    # awards: B name of leader · C max value · D how many are tied at the max
    for i, col in enumerate(["I", "J", "K", "O"]):   # champs/perfect/goose/diff
        r = 99 + i
        ws[f"A{r}"] = col
        ws[f"B{r}"] = (f"=INDEX($B$62:$B${61 + N},MATCH(MAX({col}62:"
                       f"{col}{61 + N}),{col}62:{col}{61 + N},0))")
        ws[f"C{r}"] = f"=MAX({col}62:{col}{61 + N})"
        ws[f"D{r}"] = f"=COUNTIF({col}62:{col}{61 + N},C{r})"

    ws.sheet_state = "hidden"


# ============================================================== the dashboard
def build_dashboard(ws, wb):
    thin = Side(style="thin", color=GRID)
    rule = Side(style="thin", color=BASE)
    gold_thick = Side(style="medium", color=GOLD)

    def sty(cell, *, size=10, bold=False, color=INK, fill=None, align="left",
            valign="center", fmt=None):
        c = ws[cell]
        c.font = Font(name="Arial", size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal=align, vertical=valign)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            c.number_format = fmt

    def box(rng, *, top=None, bottom=None, left=None, right=None,
            all_sides=None):
        rows = list(ws[rng])
        for r_i, row in enumerate(rows):
            for c_i, cell in enumerate(row):
                b = {"top": cell.border.top, "bottom": cell.border.bottom,
                     "left": cell.border.left, "right": cell.border.right}
                if all_sides is not None:
                    b = {k: all_sides for k in b}
                if top and r_i == 0:
                    b["top"] = top
                if bottom and r_i == len(rows) - 1:
                    b["bottom"] = bottom
                if left and c_i == 0:
                    b["left"] = left
                if right and c_i == len(row) - 1:
                    b["right"] = right
                cell.border = Border(**b)

    def fill_range(rng, color):
        for row in ws[rng]:
            for c in row:
                c.fill = PatternFill("solid", fgColor=color)

    # ---- geometry
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 1.5
    ws.column_dimensions["B"].width = 13
    for col in range(3, 22):                       # C..U
        ws.column_dimensions[get_column_letter(col)].width = 6.1
    ws.column_dimensions["V"].width = 1.5
    ws.sheet_properties.tabColor = GOLD

    for r, h in {1: 6, 2: 26, 3: 18, 4: 14, 5: 8, 6: 14, 7: 26, 8: 14, 9: 8,
                 10: 20, 11: 16, 19: 8, 20: 20, 21: 14, 29: 8, 30: 20,
                 41: 6, 42: 16, 49: 10, 50: 20, 51: 14, 52: 24, 53: 14,
                 54: 8, 55: 14}.items():
        ws.row_dimensions[r].height = h
    for r in range(12, 19):
        ws.row_dimensions[r].height = 17
    for r in range(22, 29):
        ws.row_dimensions[r].height = 14

    # ---- banner
    fill_range("B2:U4", BANNER)
    ws.merge_cells("B2:O2")
    ws["B2"] = "THE GENTLEMEN'S WAGERING COOPERATIVE"
    sty("B2", size=16, bold=True, color="FFFFFF", fill=BANNER)
    ws.merge_cells("B4:O4")
    ws["B4"] = ("OFFICIAL DASHBOARD  ·  $5 WEEKLY PARLAY  ·  "
                "CHARTER RULES IN FORCE")
    sty("B4", size=8, bold=True, color=GOLD, fill=BANNER)
    ws.merge_cells("P2:U2")
    ws["P2"] = "SEASON"
    sty("P2", size=8, bold=True, color=GOLD, fill=BANNER, align="right")
    ws.merge_cells("R3:U3")
    ws["R3"] = "2026 / 2027"
    sty("R3", size=10, bold=True, color="0000FF", fill="FFFFFF",
        align="center")
    box("R3:U3", all_sides=Side(style="thin", color=GOLD))
    dv = DataValidation(type="list", formula1='"' + ",".join(SEASONS) + '"',
                        allow_blank=False, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add("R3")
    ws.merge_cells("P4:U4")
    ws["P4"] = "pick a season — every number reworks itself"
    sty("P4", size=7, color=BASE, fill=BANNER, align="right")

    # ---- KPI tiles
    C = "'GWC Calc'!"
    tiles = [
        ("C", "F", "COOP RECORD",
         f"=SUM({C}C62:C67)&\"–\"&SUM({C}D62:D67)&\"–\"&SUM({C}E62:E67)",
         f"=SUM({C}F62:F67)&\" picks logged\"", None),
        ("H", "K", "COOP WIN RATE",
         f"=IF(SUM({C}C62:C67)+SUM({C}D62:D67)=0,0,"
         f"SUM({C}C62:C67)/(SUM({C}C62:C67)+SUM({C}D62:D67)))",
         '="pushes excluded"', "0.0%"),
        ("M", "P", "SEASON LEADER", f"={C}C70",
         f"=TEXT({C}G70,\"0.0%\")&\"  ·  \"&{C}D70&\"-\"&{C}E70"
         f"&\"  ·  \"&{C}I70&\" 🏆\"", None),
        ("R", "U", "BEST COOP WEEK",
         f"=IF({C}C60=0,\"—\",\"Week \"&{C}C60)",
         f"=IF({C}C60=0,\"\",TEXT({C}B60,\"0.0%\")&\" team win rate\")", None),
    ]
    for c1, c2, label, value_f, sub_f, fmt in tiles:
        fill_range(f"{c1}6:{c2}8", TILE)
        for rr in (6, 7, 8):
            ws.merge_cells(f"{c1}{rr}:{c2}{rr}")
        ws[f"{c1}6"] = label
        sty(f"{c1}6", size=8, bold=True, color=MUTED, fill=TILE)
        ws[f"{c1}7"] = value_f
        sty(f"{c1}7", size=17, bold=True, color=INK, fill=TILE,
            fmt=fmt or "General")
        ws[f"{c1}8"] = sub_f
        sty(f"{c1}8", size=8, color=INK2, fill=TILE)
        box(f"{c1}6:{c2}8", top=gold_thick, bottom=thin, left=thin,
            right=thin)

    # ---- standings
    ws.merge_cells("B10:U10")
    ws["B10"] = "SEASON STANDINGS"
    sty("B10", size=11, bold=True)
    box("B10:U10", bottom=rule)

    for c1, c2, t in [("B", None, "MEMBER"), ("C", None, "W"),
                      ("D", None, "L"), ("E", None, "P"), ("F", "G", "WIN %"),
                      ("H", "K", ""), ("L", "M", "LAST 4"), ("N", "O", "FORM"),
                      ("P", "Q", "CROWNS 🏆"), ("R", "S", "PERFECT 💯"),
                      ("T", "U", "0-FERS 🥚")]:
        if c2:
            ws.merge_cells(f"{c1}11:{c2}11")
        ws[f"{c1}11"] = t
        sty(f"{c1}11", size=8, bold=True, color=MUTED,
            align="left" if c1 == "B" else "center")

    medals = ["🥇 ", "🥈 ", "🥉 ", "", "", ""]
    for k in range(N):
        r, cr = 12 + k, 70 + k
        ws[f"B{r}"] = f'="{medals[k]}"&{C}C{cr}'
        sty(f"B{r}", size=10, bold=(k < 3))
        for col, src in [("C", "D"), ("D", "E"), ("E", "F")]:
            ws[f"{col}{r}"] = f"={C}{src}{cr}"
            sty(f"{col}{r}", align="center")
        ws.merge_cells(f"F{r}:G{r}")
        ws[f"F{r}"] = f"={C}G{cr}"
        sty(f"F{r}", bold=True, align="center", fmt="0.0%")
        ws.merge_cells(f"H{r}:K{r}")
        ws[f"H{r}"] = f"=REPT(\"█\",ROUND({C}H{cr}*22,0))"
        sty(f"H{r}", size=8, color=BLUE)
        ws.merge_cells(f"L{r}:M{r}")
        ws[f"L{r}"] = f'=IF({C}L{cr}="","—",{C}L{cr})'
        sty(f"L{r}", align="center", fmt="0%")
        ws.merge_cells(f"N{r}:O{r}")
        ws[f"N{r}"] = f"={C}M{cr}"
        sty(f"N{r}", size=9, bold=True, align="center")
        for col1, col2, src in [("P", "Q", "I"), ("R", "S", "J"),
                                ("T", "U", "K")]:
            ws.merge_cells(f"{col1}{r}:{col2}{r}")
            ws[f"{col1}{r}"] = f"={C}{src}{cr}"
            sty(f"{col1}{r}", align="center")
        box(f"B{r}:U{r}", bottom=thin)

    ws["B18"] = "COOP TOTAL"
    sty("B18", bold=True, color=INK2)
    for col, src in [("C", "C"), ("D", "D"), ("E", "E")]:
        ws[f"{col}18"] = f"=SUM({C}{src}62:{src}67)"
        sty(f"{col}18", bold=True, align="center", color=INK2)
    ws.merge_cells("F18:G18")
    ws["F18"] = "=IF(C18+D18=0,0,C18/(C18+D18))"
    sty("F18", bold=True, align="center", color=INK2, fmt="0.0%")
    box("B18:U18", top=rule)

    ws.conditional_formatting.add(
        "N12:O17", FormulaRule(formula=['ISNUMBER(SEARCH("HOT",N12))'],
                               font=Font(name="Arial", size=9, bold=True,
                                         color=GOOD)))
    ws.conditional_formatting.add(
        "N12:O17", FormulaRule(formula=['ISNUMBER(SEARCH("COLD",N12))'],
                               font=Font(name="Arial", size=9, bold=True,
                                         color=CRIT)))

    # ---- heatmap
    ws.merge_cells("B20:U20")
    ws["B20"] = "WEEK-BY-WEEK — MEMBER WIN % HEAT MAP"
    sty("B20", size=11, bold=True)
    box("B20:U20", bottom=rule)
    for w in range(1, WEEKS + 1):
        c = wcol(w)
        ws[f"{c}21"] = w
        sty(f"{c}21", size=8, color=MUTED, align="center")
    for i, m in enumerate(MEMBERS):
        r = 22 + i
        ws[f"B{r}"] = m
        sty(f"B{r}", size=9, color=INK2)
        for w in range(1, WEEKS + 1):
            c = wcol(w)
            ws[f"{c}{r}"] = f'=IF({C}{c}{26 + i}="","",{C}{c}{26 + i})'
            sty(f"{c}{r}", size=8, align="center", fmt="0%")
    ws["B28"] = "COOP"
    sty("B28", size=9, bold=True)
    for w in range(1, WEEKS + 1):
        c = wcol(w)
        ws[f"{c}28"] = f'=IF({C}{c}56="","",{C}{c}56)'
        sty(f"{c}28", size=8, bold=True, align="center", fmt="0%")
    box("B28:U28", top=thin)
    ws.conditional_formatting.add(
        "C22:T28",
        ColorScaleRule(start_type="num", start_value=0, start_color=RED,
                       mid_type="num", mid_value=0.5, mid_color=NEUTRAL,
                       end_type="num", end_value=1, end_color=BLUE))

    # ---- splits (left) & chart (right)
    ws.merge_cells("B30:H30")
    ws["B30"] = "HOW THE COOP WINS"
    sty("B30", size=11, bold=True)
    box("B30:H30", bottom=rule)
    ws.merge_cells("J30:Q30")
    ws["J30"] = ('=IF(\'GWC Live\'!$A$1="","THIS WEEK — NO SLATE LOADED YET",'
                 '"THIS WEEK — LIVE SLATE · WEEK "&\'GWC Live\'!$A$1)')
    sty("J30", size=11, bold=True)
    ws.merge_cells("R30:U30")
    ws["R30"] = '=IF(\'GWC Live\'!$B$1="","","as of "&\'GWC Live\'!$B$1)'
    sty("R30", size=7.5, color=MUTED, align="right")
    box("J30:U30", bottom=rule)

    for col, t in [("C", "W"), ("D", "L"), ("E", "WIN %")]:
        ws[f"{col}31"] = t
        sty(f"{col}31", size=8, bold=True, color=MUTED, align="center")

    def split_row(r, label, cr):
        ws[f"B{r}"] = label
        sty(f"B{r}", size=9, color=INK2)
        for col, src in [("C", "C"), ("D", "D")]:
            ws[f"{col}{r}"] = f"={C}{src}{cr}"
            sty(f"{col}{r}", size=9, align="center")
        ws[f"E{r}"] = f"={C}E{cr}"
        sty(f"E{r}", size=9, bold=True, align="center", fmt="0%")
        ws.merge_cells(f"F{r}:H{r}")
        ws[f"F{r}"] = f"=REPT(\"█\",ROUND({C}E{cr}*14,0))"
        sty(f"F{r}", size=8, color=BLUE)
        box(f"B{r}:H{r}", bottom=thin)

    labels = ["Spreads", "Over/Unders", "Moneylines", "Favorites",
              "Underdogs", "Home teams", "Away teams", "Overs", "Unders"]
    for i, label in enumerate(labels):
        split_row(32 + i, label, 78 + i)

    ws.merge_cells("B42:H42")
    ws["B42"] = "BY DAY"
    sty("B42", size=9, bold=True, color=MUTED)
    for i, d in enumerate(["Thursday", "Friday", "Saturday", "Sunday",
                           "Monday"]):
        split_row(43 + i, d, 92 + i)

    # ---- live-week panel (fed by the hidden 'GWC Live' sheet)
    L = "'GWC Live'!"
    for c1, c2, t in [("J", None, "DAY"), ("K", "L", "MEMBER"),
                      ("M", "N", "MATCHUP"), ("O", "Q", "PICK"),
                      ("R", "T", "SCORE / KICKOFF"), ("U", None, "RES")]:
        if c2:
            ws.merge_cells(f"{c1}31:{c2}31")
        ws[f"{c1}31"] = t
        sty(f"{c1}31", size=8, bold=True, color=MUTED,
            align="left" if c1 in ("K", "M", "O", "R") else "center")
    for i in range(16):                      # up to 16 games in an NFL week
        r, fr = 32 + i, 4 + i
        ws.row_dimensions[r].height = 15
        ws[f"J{r}"] = f'=IF({L}$B{fr}="","",{L}$A{fr})'
        sty(f"J{r}", size=9, color=INK2, align="center")
        ws.merge_cells(f"K{r}:L{r}")
        ws[f"K{r}"] = f'=IF({L}$B{fr}="","",{L}$B{fr})'
        sty(f"K{r}", size=9, bold=True)
        ws.merge_cells(f"M{r}:N{r}")
        ws[f"M{r}"] = f'=IF({L}$B{fr}="","",{L}$C{fr})'
        sty(f"M{r}", size=9, color=INK2)
        ws.merge_cells(f"O{r}:Q{r}")
        ws[f"O{r}"] = f'=IF({L}$B{fr}="","",{L}$D{fr})'
        sty(f"O{r}", size=9, bold=True)
        ws.merge_cells(f"R{r}:T{r}")
        ws[f"R{r}"] = f'=IF({L}$B{fr}="","",{L}$E{fr})'
        sty(f"R{r}", size=9, color=INK2)
        ws[f"U{r}"] = f'=IF({L}$F{fr}="","",{L}$F{fr})'
        sty(f"U{r}", size=9, bold=True, align="center")
        box(f"J{r}:U{r}", bottom=thin)
    ws.conditional_formatting.add(
        "U32:U47", FormulaRule(formula=['U32="W"'],
                               font=Font(name="Arial", size=9, bold=True,
                                         color=GOOD)))
    ws.conditional_formatting.add(
        "U32:U47", FormulaRule(formula=['U32="L"'],
                               font=Font(name="Arial", size=9, bold=True,
                                         color=CRIT)))
    ws.conditional_formatting.add(
        "U32:U47", FormulaRule(formula=['U32="P"'],
                               font=Font(name="Arial", size=9, bold=True,
                                         color=MUTED)))

    ws.merge_cells("B55:U55")
    ws["B55"] = (f'="Live off the Games raw table · data through Week "'
                 f"&{C}B59&\" · outcomes are uncertain, losses are shared, "
                 f'and results speak for themselves — Charter §I"')
    sty("B55", size=7.5, color=MUTED)


def main(path):
    src = Path(path).expanduser()
    if not src.exists():
        sys.exit(f"Workbook not found: {src}")
    stamp = datetime.now().strftime("%Y-%m-%d %H%M%S")
    backup = src.with_name(f"{src.stem} (backup {stamp}){src.suffix}")
    shutil.copy2(src, backup)
    print(f"backup: {backup.name}")

    wb = openpyxl.load_workbook(src)
    for name in ("Dashboard", "GWC Calc"):
        if name in wb.sheetnames:
            del wb[name]
    if "GWC Live" not in wb.sheetnames:
        live = wb.create_sheet("GWC Live")
        live.sheet_state = "hidden"
    calc = wb.create_sheet("GWC Calc")
    build_calc(calc)
    dash = wb.create_sheet("Dashboard", 0)
    build_dashboard(dash, wb)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(src)
    print(f"dashboard built into: {src}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
