"""One-time import of past-season picks from the Commissioner's Excel workbook.

Usage: python3 seed_history.py "path/to/Gentlemens Wagering Cooperative.xlsx"

Reads the raw data table on the 'Games' sheet and loads weeks, games, and
graded picks into the app database so standings and analytics carry over.
Safe to re-run: it skips seasons that are already in the database.
"""
import sys

import openpyxl
from sqlalchemy import select

from gwc.db import assignments, engine, games, players, rows, weeks


def parse_score(text, away, home):
    # "Bears 19 - Vikings 17" -> (19, 17) as (away, home)
    try:
        left, right = text.split(" - ")
        a_name, a_pts = left.rsplit(" ", 1)
        h_name, h_pts = right.rsplit(" ", 1)
        if a_name.strip() == home:  # score listed home-first
            return int(h_pts), int(a_pts)
        return int(a_pts), int(h_pts)
    except Exception:
        return None, None


def main(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Games"]

    # find the raw-data header row (Season | Week | Day | Person | ...)
    header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row[0] == "Season" and row[1] == "Week":
            header_row = i
            break
    if not header_row:
        sys.exit("Could not find the raw data table on the 'Games' sheet.")

    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row[0] or not row[3] or row[0] == "Season":
            continue
        records.append({
            "season": str(row[0]).strip(), "week": int(row[1]),
            "day": str(row[2]).strip(), "person": str(row[3]).strip(),
            "matchup": str(row[4]).strip(), "pick": str(row[5]).strip(),
            "score": str(row[6]).strip() if row[6] else None,
            "result": str(row[7]).strip() if row[7] else None,
            "pick_type": str(row[8]).strip() if row[8] else None,
            "fav_dog": str(row[9]).strip() if row[9] not in (None, 0, "0") else None,
            "home_away": str(row[10]).strip() if row[10] not in (None, 0, "0") else None,
        })
    print(f"Found {len(records)} historical picks.")

    eng = engine()
    with eng.begin() as conn:
        existing_seasons = {w["season"] for w in rows(conn, select(weeks))}
        pids = {p["name"]: p["id"] for p in rows(conn, select(players))}

        week_ids, game_ids = {}, {}
        imported = 0
        for r in records:
            if r["season"] in existing_seasons:
                continue
            wkey = (r["season"], r["week"])
            if wkey not in week_ids:
                week_ids[wkey] = conn.execute(weeks.insert().values(
                    season=r["season"], week_num=r["week"], status="final",
                )).inserted_primary_key[0]

            sep = " @ " if " @ " in r["matchup"] else " vs "
            away, home = [t.strip() for t in r["matchup"].split(sep, 1)]
            gkey = (*wkey, r["matchup"])
            if gkey not in game_ids:
                a_s, h_s = (parse_score(r["score"], away, home)
                            if r["score"] else (None, None))
                game_ids[gkey] = conn.execute(games.insert().values(
                    week_id=week_ids[wkey], day=r["day"], away=away, home=home,
                    away_score=a_s, home_score=h_s, final=True,
                )).inserted_primary_key[0]

            pid = pids.get(r["person"])
            if not pid:
                print(f"  ! unknown member {r['person']!r}, skipped")
                continue
            conn.execute(assignments.insert().values(
                week_id=week_ids[wkey], game_id=game_ids[gkey], player_id=pid,
                pick_type=r["pick_type"], pick_selection=r["pick"],
                fav_dog=r["fav_dog"], home_away=r["home_away"],
                result=r["result"],
            ))
            imported += 1
    print(f"Imported {imported} picks across {len(week_ids)} weeks.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
