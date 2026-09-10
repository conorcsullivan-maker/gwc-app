"""Excel bridge: sync picks into the Commissioner's original workbook, and
build a standalone multi-sheet report snapshot for the group chat."""
from __future__ import annotations
import io
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from sqlalchemy import select

from . import schedule, stats
from .db import (CURRENT_SEASON, MEMBERS, SEASON_YEAR, assignments, engine,
                 games, players, rows, weeks)

RAW_HEADER = ["Season", "Week", "Day", "Person", "Matchup", "Pick", "Score",
              "Result", "Pick Type", "Fav/Dog", "Home/Away", "Spread #",
              "Game Key"]


def _all_graded_rows():
    stmt = (select(assignments, players.c.name.label("player"),
                   weeks.c.season, weeks.c.week_num,
                   games.c.day, games.c.away, games.c.home,
                   games.c.away_score, games.c.home_score)
            .join(players, assignments.c.player_id == players.c.id)
            .join(weeks, assignments.c.week_id == weeks.c.id)
            .join(games, assignments.c.game_id == games.c.id)
            .where(assignments.c.result.isnot(None))
            .order_by(weeks.c.season, weeks.c.week_num))
    with engine().connect() as conn:
        return rows(conn, stmt)


def _to_sheet_row(r):
    matchup = f"{r['away']} @ {r['home']}"
    score = ""
    if r["away_score"] is not None:
        score = (f"{r['away']} {r['away_score']} - "
                 f"{r['home']} {r['home_score']}")
    is_total = r["pick_type"] == "Over/Under"
    spread_num = (abs(r["pick_line"])
                  if r["pick_type"] == "Spread" and r["pick_line"] is not None
                  else "")
    return [r["season"], r["week_num"], r["day"] or "", r["player"], matchup,
            r["pick_selection"] or "", score, r["result"],
            r["pick_type"] or "",
            "0" if is_total else (r["fav_dog"] or ""),   # workbook convention
            "0" if is_total else (r["home_away"] or ""),
            spread_num, f"{r['week_num']}|{matchup}"]


def sync_workbook(xlsx_path: str) -> tuple[int, str]:
    """Append graded picks missing from the workbook's Games raw table.

    A timestamped backup copy is written next to the original first.
    Returns (rows_added, backup_path).
    """
    path = Path(xlsx_path).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")
    stamp = datetime.now().strftime("%Y-%m-%d %H%M%S")
    backup = path.with_name(f"{path.stem} (backup {stamp}){path.suffix}")
    shutil.copy2(path, backup)

    wb = openpyxl.load_workbook(path)
    ws = wb["Games"]

    header_row = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row[0] == "Season" and row[1] == "Week":
            header_row = i
            break
    if header_row is None:
        raise ValueError("Couldn't find the raw data table (Season | Week | "
                         "...) on the 'Games' sheet.")

    existing = set()
    last_row = header_row
    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1,
                                         values_only=True),
                            start=header_row + 1):
        if row[0] and row[3]:
            # neutral-site games appear as "X vs Y" in the workbook
            matchup = str(row[4]).strip().replace(" vs ", " @ ")
            existing.add((str(row[0]).strip(), int(row[1]),
                          str(row[3]).strip(), matchup))
            last_row = i

    added = 0
    for r in _all_graded_rows():
        sheet_row = _to_sheet_row(r)
        key = (sheet_row[0], int(sheet_row[1]), sheet_row[3], sheet_row[4])
        if key in existing:
            continue
        last_row += 1
        for col, val in enumerate(sheet_row, start=1):
            ws.cell(row=last_row, column=col, value=val)
        existing.add(key)
        added += 1

    if added:
        wb.save(path)
    return added, str(backup)


def write_live_feed(xlsx_path: str) -> int:
    """Refresh the hidden 'GWC Live' sheet the dashboard's THIS WEEK panel
    reads: one row per assignment of the latest current-season week with
    member, matchup, pick, live score/kickoff, and result.

    Returns the number of rows written (0 if no current-season week exists).
    """
    from datetime import datetime as dt
    from zoneinfo import ZoneInfo

    path = Path(xlsx_path).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    with engine().connect() as conn:
        wks = rows(conn, select(weeks).where(weeks.c.season == CURRENT_SEASON)
                   .order_by(weeks.c.week_num.desc()))
        week = wks[0] if wks else None
        alist = []
        if week:
            alist = rows(conn, select(
                assignments.c.pick_selection, assignments.c.result,
                assignments.c.pick_type, assignments.c.fav_dog,
                assignments.c.home_away,
                players.c.name.label("player"), games.c.day,
                games.c.kickoff_et, games.c.away, games.c.home,
                games.c.away_abbr, games.c.home_abbr, games.c.espn_id)
                .select_from(assignments
                             .join(players, assignments.c.player_id == players.c.id)
                             .join(games, assignments.c.game_id == games.c.id))
                .where(assignments.c.week_id == week["id"])
                .order_by(games.c.kickoff_et, players.c.name))

    live = {}
    if week:
        try:
            live = schedule.fetch_week_live(SEASON_YEAR, week["week_num"])
        except Exception:
            live = {}          # offline: kickoff times still render

    wb = openpyxl.load_workbook(path)
    ws = (wb["GWC Live"] if "GWC Live" in wb.sheetnames
          else wb.create_sheet("GWC Live"))
    ws.sheet_state = "hidden"
    for row in ws["A1:H40"]:
        for c in row:
            c.value = None

    if week:
        ws["A1"] = week["week_num"]
        et = ZoneInfo("America/New_York")
        ws["B1"] = dt.now(et).strftime("%b %-d, %-I:%M %p ET")
        for i, a in enumerate(alist):
            ko = dt.fromisoformat(a["kickoff_et"]) if a["kickoff_et"] else None
            g = live.get(a["espn_id"], {})
            state = g.get("state", "pre")
            if state == "pre" or not g:
                score = ko.strftime("%a %-m/%-d %-I:%M %p") if ko else "TBD"
            else:
                score = (f"{a['away_abbr']} {g['away_score']} – "
                         f"{a['home_abbr']} {g['home_score']} · {g['detail']}")
            r = 4 + i
            ws[f"A{r}"] = a["day"] or (ko.strftime("%a") if ko else "")
            ws[f"B{r}"] = a["player"]
            ws[f"C{r}"] = f"{a['away_abbr']} @ {a['home_abbr']}"
            ws[f"D{r}"] = a["pick_selection"] or "—"
            ws[f"E{r}"] = score
            ws[f"F{r}"] = {"Win": "W", "Loss": "L", "Push": "P"}.get(
                a["result"], "")

    # also keep the human-facing 'Current Week' sheet in step
    if week and "Current Week" in wb.sheetnames:
        cw = wb["Current Week"]
        for row in cw["A4:K30"]:
            for c in row:
                c.value = None
        for i, a in enumerate(alist[:26]):
            g = live.get(a["espn_id"], {})
            score = ""
            if g.get("state") == "post":
                score = (f"{a['away']} {g['away_score']} - "
                         f"{a['home']} {g['home_score']}")
            r = 4 + i
            cw[f"A{r}"] = CURRENT_SEASON
            cw[f"B{r}"] = week["week_num"]
            cw[f"C{r}"] = a["day"]
            cw[f"D{r}"] = a["player"]
            cw[f"E{r}"] = f"{a['away']} @ {a['home']}"
            cw[f"F{r}"] = a["pick_selection"] or ""
            cw[f"G{r}"] = score
            cw[f"H{r}"] = a["result"] or ""
            cw[f"I{r}"] = a["pick_type"] or ""
            cw[f"J{r}"] = a["fav_dog"] or ""
            cw[f"K{r}"] = a["home_away"] or ""
        for dv in cw.data_validations.dataValidation:
            if "Jake" in (dv.formula1 or ""):    # stale roster in dropdown
                dv.formula1 = '"' + ",".join(MEMBERS) + '"'
    wb.save(path)
    return len(alist)


def snapshot_bytes() -> bytes:
    """A standalone .xlsx report: raw picks, standings, weekly champions,
    analytics splits, and bragging rights — good for sharing with the group."""
    graded = _all_graded_rows()
    picks = stats.enrich([{
        "player": r["player"], "week_num": r["week_num"], "season": r["season"],
        "result": r["result"], "pick_type": r["pick_type"],
        "pick_selection": r["pick_selection"], "fav_dog": r["fav_dog"],
        "home_away": r["home_away"], "day": r["day"], "away": r["away"],
        "home": r["home"], "pick_team": r["pick_team"],
        "pick_line": r["pick_line"],
    } for r in graded])

    wb = openpyxl.Workbook()
    bold = Font(name="Arial", bold=True)
    plain = Font(name="Arial")

    def put(ws, row, values, header=False):
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = bold if header else plain
        return row + 1

    def put_table(ws, row, title, dict_rows, pct_keys=("Win %",)):
        row = put(ws, row, [title], header=True)
        if not dict_rows:
            return put(ws, row, ["(no data)"]) + 1
        keys = list(dict_rows[0].keys())
        row = put(ws, row, keys, header=True)
        for d in dict_rows:
            vals = [round(d[k], 3) if k in pct_keys and isinstance(d[k], float)
                    else d[k] for k in keys]
            row = put(ws, row, vals)
        return row + 1

    ws = wb.active
    ws.title = "Picks"
    r = put(ws, 1, RAW_HEADER, header=True)
    for g in graded:
        r = put(ws, r, _to_sheet_row(g))

    seasons = sorted({p["season"] for p in picks})
    ws = wb.create_sheet("Standings")
    r = 1
    for season in seasons:
        sp = [p for p in picks if p["season"] == season]
        table, coop = stats.standings(sp, MEMBERS)
        r = put_table(ws, r, f"SEASON {season}", table)
        r = put(ws, r, ["Coop total", coop["Wins"], coop["Losses"],
                        coop["Pushes"], coop["Total"],
                        round(coop["Win %"], 3)]) + 1

    ws = wb.create_sheet("Weekly Report")
    r = 1
    for season in seasons:
        sp = [p for p in picks if p["season"] == season]
        r = put_table(ws, r, f"WEEKLY CHAMPIONS — {season}",
                      stats.weekly_report(sp), pct_keys=("Coop Win %",))

    ws = wb.create_sheet("Analytics")
    overview, per_player = stats.analytics(picks, MEMBERS)
    r = put_table(ws, 1, "GROUP SPLITS (ALL TIME)", overview)
    r = put_table(ws, r, "PER PLAYER", per_player)
    r = put_table(ws, r, "BY DAY", stats.day_splits(picks))
    r = put_table(ws, r, "SPREAD SIZE", stats.spread_buckets(picks))
    r = put_table(ws, r, "SITUATIONAL MATRIX",
                  stats.player_matrix(picks, MEMBERS), pct_keys=())

    ws = wb.create_sheet("Bragging Rights")
    awards, perfect, goose = stats.bragging_rights(picks, MEMBERS)
    r = put_table(ws, 1, "AWARDS", awards, pct_keys=())
    r = put_table(ws, r, "PERFECT WEEKS", perfect, pct_keys=())
    r = put_table(ws, r, "GOOSE EGGS", goose, pct_keys=())
    put_table(ws, r, "PICK STREAKS", stats.pick_streaks(picks, MEMBERS),
              pct_keys=())

    for ws in wb.worksheets:
        for col_cells in ws.columns:
            width = max((len(str(c.value)) for c in col_cells if c.value),
                        default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(
                width + 2, 34)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
